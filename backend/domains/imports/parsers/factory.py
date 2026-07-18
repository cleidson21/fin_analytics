from typing import List, Union, Dict, Any
from pydantic import BaseModel
import csv
import io
import hashlib
from datetime import datetime, timezone
import logging
import openpyxl

from backend.domains.imports.dto import RawDocument, ImportSource
from backend.domains.imports.readers.pierre_txt import PierreTxtReader
from backend.domains.imports.parsers.pierre_txt import PierreTxtParser
from backend.domains.imports.parsers.nubank_conta import NubankContaParser
from backend.domains.imports.parsers.nubank_fatura import NubankFaturaParser
from backend.domains.imports.parsers.myprofit_proventos import MyProfitProventosParser

logger = logging.getLogger(__name__)

# --- New specific parsers for Evolucao and B3 ---

class MyProfitEvolucaoParser:
    def parse(self, doc: RawDocument) -> List[Any]:
        parsed_rows = []
        for row in doc.content:
            if not row.get("Data"): continue
            
            def clean_num(val):
                if not val: return 0.0
                val = str(val).replace("R$", "").replace("%", "").strip()
                val = val.replace(".", "").replace(",", ".")
                try: return float(val)
                except: return 0.0

            parsed_rows.append(type("EvolucaoRow", (), {
                "model_dump_json": lambda self, r=row: __import__('json').dumps({
                    "reference_month": r.get("Data"),
                    "patrimonio": clean_num(r.get("Patrimônio")),
                    "compras": clean_num(r.get("Compras")),
                    "vendas": clean_num(r.get("Vendas"))
                })
            })())
        return parsed_rows

class B3Parser:
    def parse(self, doc: RawDocument) -> List[Any]:
        parsed_rows = []
        for row in doc.content:
            parsed_rows.append(type("B3Row", (), {
                "model_dump_json": lambda self, r=row: __import__('json').dumps({
                    "ativo": r.get("Produto", "UNKNOWN"),
                    "quantidade": r.get("Quantidade", 0),
                    "valor": r.get("Valor Operação", 0)
                })
            })())
        return parsed_rows

# --- Factory ---

class ImportParserFactory:
    @staticmethod
    def _read_csv(raw_content: Union[str, bytes], filename: str, source: ImportSource) -> RawDocument:
        if isinstance(raw_content, bytes):
            try:
                raw_text = raw_content.decode('utf-8-sig')
            except UnicodeDecodeError:
                raw_text = raw_content.decode('latin-1', errors='replace')
        else:
            raw_text = raw_content
            
        f = io.StringIO(raw_text)
        # Try to detect delimiter (semicolon or comma)
        # B3 usually uses ';' while Nubank uses ','
        first_line = raw_text.split('\\n')[0] if '\\n' in raw_text else raw_text
        delimiter = ';' if ';' in first_line else ','
        
        reader = csv.DictReader(f, delimiter=delimiter)
        content_dict = [row for row in reader]
        
        doc_hash = hashlib.sha256(raw_text.encode('utf-8')).hexdigest()
        
        return RawDocument(
            source=source,
            filename=filename,
            mime_type="text/csv",
            encoding="utf-8",
            parser_version="1.0.0",
            metadata={"total_lines": len(content_dict)},
            content=content_dict,
            document_hash=doc_hash,
            created_at=datetime.now(timezone.utc)
        )

    @staticmethod
    def _read_xlsx(raw_content: bytes, filename: str, source: ImportSource) -> RawDocument:
        wb = openpyxl.load_workbook(io.BytesIO(raw_content), data_only=True)
        ws = wb.active
        
        content_dict = []
        headers = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(cell) if cell else f"col_{j}" for j, cell in enumerate(row)]
            else:
                row_dict = {}
                for j, cell in enumerate(row):
                    if j < len(headers):
                        row_dict[headers[j]] = cell
                content_dict.append(row_dict)
                
        doc_hash = hashlib.sha256(raw_content).hexdigest()
        
        return RawDocument(
            source=source,
            filename=filename,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            encoding="binary",
            parser_version="1.0.0",
            metadata={"total_lines": len(content_dict)},
            content=content_dict,
            document_hash=doc_hash,
            created_at=datetime.now(timezone.utc)
        )

    @classmethod
    def execute(cls, raw_content: Union[str, bytes], filename: str, source_type: str) -> List[Any]:
        source = getattr(ImportSource, source_type, ImportSource.NUBANK_CONTA)
        
        if source == ImportSource.PIERRE:
            raw_text = raw_content.decode('utf-8', errors='replace') if isinstance(raw_content, bytes) else raw_content
            lines = raw_text.splitlines()
            doc = RawDocument(
                source=ImportSource.PIERRE,
                filename=filename,
                mime_type="text/plain",
                encoding="utf-8",
                parser_version="1.0",
                content=lines,
                document_hash="hash",
                created_at=datetime.now(timezone.utc)
            )
            parser = PierreTxtParser()
            return parser.parse(doc)
            
        elif source == ImportSource.B3:
            # Excel ou CSV?
            if filename.lower().endswith(".xlsx"):
                if isinstance(raw_content, str):
                    raise ValueError("Arquivo XLSX deve ser recebido em formato binário.")
                doc = cls._read_xlsx(raw_content, filename, source)
            else:
                doc = cls._read_csv(raw_content, filename, source)
            parser = B3Parser()
            return parser.parse(doc)
            
        else:
            # Padrão CSV
            doc = cls._read_csv(raw_content, filename, source)
            
            if source == ImportSource.NUBANK_CONTA:
                return NubankContaParser().parse(doc)
            elif source == ImportSource.NUBANK_FATURA:
                return NubankFaturaParser().parse(doc)
            elif source == ImportSource.MYPROFIT_PROVENTOS:
                return MyProfitProventosParser().parse(doc)
            elif source == ImportSource.MYPROFIT_EVOLUCAO:
                return MyProfitEvolucaoParser().parse(doc)
            
        return []
